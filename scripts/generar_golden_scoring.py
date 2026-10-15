"""Genera tests/golden/scoring_golden.json: casos de respuestas con el resultado que calculan
las calculadoras de Excel originales (docs/referencia), recalculadas con LibreOffice.

Uso:  python scripts/generar_golden_scoring.py        (requiere `soffice` en el PATH)

tests/test_scoring_golden.py compara modules.scoring contra ese archivo, así que la calificación
queda verificada contra las fórmulas reales sin necesitar LibreOffice al correr pytest.
"""
import json
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
REF = RAIZ / "docs" / "referencia"
SALIDA = RAIZ / "tests" / "golden" / "scoring_golden.json"
CODIGOS = {"Mercadotecnia": "MKT", "Capital Humano": "CH", "Administración Financiera": "AF",
           "Negocios Internacionales": "NI", "Emprendimiento": "EMP"}
DIMS = {"Apertura a la Experiencia": "APERTURA", "Responsabilidad / Tesón": "RESPONSABILIDAD",
        "Extraversión": "EXTRAVERSION", "Amabilidad / Cordialidad": "AMABILIDAD",
        "Estabilidad Emocional": "ESTABILIDAD"}


def casos_test1(rng):
    ids = [f"BF{n:02d}" for n in range(1, 26)]
    casos = [[1] * 25, [5] * 25, [3] * 25, [1] * 5 + [5] * 5 + [3] * 15, [5] * 5 + [1] * 20]
    casos += [[rng.randint(1, 5) for _ in ids] for _ in range(60)]
    # rangos estrechos para provocar empates y valores justo en los umbrales (18, 14, 0.76, 0.56)
    casos += [[rng.choice([3, 4]) for _ in ids] for _ in range(20)]
    casos += [[rng.choice([2, 3, 4, 5]) for _ in ids] for _ in range(20)]
    return [dict(zip(ids, c)) for c in casos]


def casos_test2(rng):
    ids = ["M01", "M02", "M04", "M05", "M09", "CH01", "CH03", "CH04", "CH06", "CH10", "AF01", "AF02",
           "AF03", "AF04", "AF05", "NI01", "NI04", "NI05", "NI06", "NI08", "E01", "E03", "E04", "E05", "E10"]
    casos = [[0] * 25, [4] * 25, [2] * 25]
    casos += [[rng.randint(0, 4) for _ in ids] for _ in range(50)]
    casos += [[rng.choice([2, 3, 4]) for _ in ids] for _ in range(30)]
    con_se = []
    for _ in range(60):  # "SE" con probabilidad creciente
        p = rng.choice([0.1, 0.2, 0.35])
        con_se.append([("SE" if rng.random() < p else rng.randint(0, 4)) for _ in ids])
    casos += con_se
    base = [3] * 25
    casos += [base[:5] + [3] * 20 for _ in range(1)]
    solo1 = list(base); solo1[0] = "SE"; casos.append(solo1)                       # 1 SE: sin exploración
    dos = list(base); dos[0] = dos[1] = "SE"; casos.append(dos)                     # 2 SE: exploración
    todas = list(base); todas[0:5] = ["SE"] * 5; casos.append(todas)               # una especialidad completa en SE
    todas2 = [rng.randint(0, 4) for _ in ids]; todas2[10:15] = ["SE"] * 5; casos.append(todas2)
    return [dict(zip(ids, c)) for c in casos]


def preparar(origen, hoja, celdas_id, col_resp, casos, destino, prefijo):
    archivos = []
    for i, caso in enumerate(casos):
        wb = openpyxl.load_workbook(origen)
        ws = wb[hoja]
        for fila, pid in celdas_id(ws):
            ws.cell(fila, col_resp).value = caso[pid]
        ruta = destino / f"{prefijo}_{i:03d}.xlsx"
        wb.save(ruta)
        archivos.append(ruta)
    return archivos


def convertir(archivos, sal, lote=30):
    """Recalcula con LibreOffice por lotes y reintenta los que no salieron."""
    pendientes = list(archivos)
    for intento in range(3):
        for i in range(0, len(pendientes), lote):
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(sal),
                            *map(str, pendientes[i:i + lote])], check=True, capture_output=True, timeout=600)
        pendientes = [f for f in pendientes if not (sal / f.name).exists()]
        if not pendientes:
            return
    raise RuntimeError(f"LibreOffice no convirtió: {[f.name for f in pendientes][:5]}")


def main():
    rng = random.Random(20261014)
    c1, c2 = casos_test1(rng), casos_test2(rng)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        ent, sal = tmp / "ent", tmp / "sal"
        ent.mkdir(); sal.mkdir()
        f1 = preparar(REF / "1_test_big_five_fca_calculadora.xlsx", "Test y Respuestas",
                      lambda ws: [(r, f"BF{int(ws.cell(r, 1).value):02d}") for r in range(9, 34)], 4, c1, ent, "t1")
        f2 = preparar(REF / "2_test_orientacion_especialidades_administracion.xlsx", "Test_Orientación",
                      lambda ws: [(r, ws.cell(r, 3).value) for r in range(23, 48)], 5, c2, ent, "t2")
        convertir(f1 + f2, sal)
        out = {"test1": [], "test2": []}
        for caso, f in zip(c1, f1):
            wb = openpyxl.load_workbook(sal / f.name, data_only=True)
            r = wb["Resultados y Diagnóstico"]
            out["test1"].append({
                "respuestas": caso,
                "dimensiones": {DIMS[r.cell(x, 1).value]: {"obtenido": r.cell(x, 2).value, "nivel": r.cell(x, 4).value}
                                for x in range(5, 10)},
                "afinidades": {CODIGOS[r.cell(x, 1).value]: {"porcentaje": r.cell(x, 3).value, "nivel": r.cell(x, 4).value}
                               for x in range(14, 19)},
                "principal": CODIGOS[r["A30"].value.replace(" (Especialidad Principal)", "")],
            })
        for caso, f in zip(c2, f2):
            wb = openpyxl.load_workbook(sal / f.name, data_only=True)
            r = wb["Test_Orientación"]
            out["test2"].append({
                "respuestas": caso,
                "especialidades": {CODIGOS[r.cell(x, 1).value]: {
                    "obtenido": r.cell(x, 2).value, "maximo": r.cell(x, 3).value, "porcentaje": r.cell(x, 4).value,
                    "nivel": r.cell(x, 5).value, "requiere_exploracion": "Requiere" in str(r.cell(x, 6).value)}
                    for x in range(51, 56)},
                "principal": CODIGOS[r["B59"].value], "secundaria": CODIGOS[r["B60"].value],
                "empate_tecnico": str(r["B74"].value).startswith("PERFIL HÍBRIDO"),
            })
    SALIDA.parent.mkdir(exist_ok=True)
    SALIDA.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(out['test1'])} casos test 1, {len(out['test2'])} casos test 2 -> {SALIDA}")


if __name__ == "__main__":
    main()
