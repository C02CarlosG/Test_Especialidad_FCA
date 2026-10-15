"""Convierte las calculadoras de Excel de docs/referencia/ al formato del cargador
(config/tests/*.xlsx: hojas especialidades, preguntas, escala, info).

Uso:  python scripts/convertir_tests_reales.py

Solo copia texto de las calculadoras (reactivos, escalas, descripciones); no inventa contenido.
Lo que el cargador no representa (combinaciones OCEAN -> especialidad del test 1, opción
scoring de "SE" del test 2, textos por nivel) se queda en las calculadoras para la Fase 4.
"""
from math import ceil
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent.parent
REF = RAIZ / "docs" / "referencia"
DESTINO = RAIZ / "config" / "tests"
POR_BLOQUE = 10


def escribir(ruta: Path, titulo, especialidades, preguntas, escala) -> None:
    hojas = {
        "info": pd.DataFrame([("titulo", titulo)], columns=["clave", "valor"]),
        "especialidades": pd.DataFrame(especialidades, columns=["codigo", "nombre", "descripcion", "perfil_profesional"]),
        "preguntas": pd.DataFrame(
            [(pid, ceil(n / POR_BLOQUE), txt, esp, "NO") for n, (pid, txt, esp) in enumerate(preguntas, 1)],
            columns=["id", "bloque", "texto", "especialidad_codigo", "invertida"]),
        "escala": pd.DataFrame(escala, columns=["valor", "etiqueta"]),
    }
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            df.to_excel(w, sheet_name=nombre, index=False)


def test_1() -> None:
    ws = openpyxl.load_workbook(REF / "1_test_big_five_fca_calculadora.xlsx")["Test y Respuestas"]
    dims = {  # nombre en la calculadora -> código
        "Apertura a la Experiencia": "APERTURA", "Responsabilidad / Tesón": "RESPONSABILIDAD",
        "Extraversión": "EXTRAVERSION", "Amabilidad / Cordialidad": "AMABILIDAD",
        "Estabilidad Emocional": "ESTABILIDAD",
    }
    preguntas = [(f"BF{int(ws.cell(r, 1).value):02d}", ws.cell(r, 3).value.strip(), dims[ws.cell(r, 2).value])
                 for r in range(9, 34)]
    # Las 5 "especialidades" de este test son las dimensiones OCEAN; su relación con las
    # especialidades de la carrera (p. ej. Mercadotecnia = Apertura + Extraversión) se calcula en la Fase 4.
    esp = [(c, n, "", "") for n, c in dims.items()]
    escala = []
    for parte in ws["A6"].value.split("|"):
        valor, etiqueta = parte.split("=")
        escala.append((int(valor), etiqueta.strip()))
    escribir(DESTINO / "1_test_big_five.xlsx", "Big Five (OCEAN)", esp, preguntas, escala)


def test_2() -> None:
    wb = openpyxl.load_workbook(REF / "2_test_orientacion_especialidades_administracion.xlsx")
    ws, guia = wb["Test_Orientación"], wb["Guía_Tutoría"]
    codigos = {"Mercadotecnia": "MKT", "Capital Humano": "CH", "Administración Financiera": "AF",
               "Negocios Internacionales": "NI", "Emprendimiento": "EMP"}
    preguntas = [(ws.cell(r, 3).value.strip(), ws.cell(r, 4).value.strip(), codigos[ws.cell(r, 2).value])
                 for r in range(23, 48)]
    esp = []
    for r in range(3, 8):
        nombre = guia.cell(r, 1).value
        perfil = f"DISC: {guia.cell(r, 2).value} Holland: {guia.cell(r, 3).value} HBDI: {guia.cell(r, 4).value}"
        esp.append((codigos[nombre], nombre, guia.cell(r, 5).value, perfil))
    # Escala 0-4 más la fila especial "SE" (sin experiencia: dato faltante, no es 0).
    escala = [(int(ws.cell(r, 1).value), ws.cell(r, 2).value) for r in range(14, 19)]
    escala.append(("SE", ws.cell(19, 2).value))
    escribir(DESTINO / "2_test_orientacion_especialidades.xlsx",
             "Orientación para especialidades en Administración", esp, preguntas, escala)


if __name__ == "__main__":
    test_1()
    test_2()
    print("Listo:", *sorted(p.name for p in DESTINO.glob("*.xlsx")), sep="\n- ")
