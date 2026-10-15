import io
import zipfile
from dataclasses import replace
from datetime import datetime, timezone

import openpyxl
import pytest

from modules.acceso import Alumno
from modules.panel import PanelDatos, cargar_panel
from modules.reports_admin import excel_concentrado, nombre_pdf, zip_pdfs
from test_panel import AHORA, TEST1, TEST2, poblar

FECHA = datetime(2026, 10, 15, 10, 0, tzinfo=timezone.utc)


@pytest.fixture(scope="module")
def panel():
    return cargar_panel(poblar(), (TEST1, TEST2))


def _libro(panel_):
    return openpyxl.load_workbook(io.BytesIO(excel_concentrado(panel_, FECHA)))


def test_hojas_del_concentrado(panel):
    assert _libro(panel).sheetnames == ["Resumen", "Alumnos", "Test 1", "Test 2", "Respuestas"]


def test_resumen_cuenta_alumnos_y_primeras_opciones(panel):
    ws = _libro(panel)["Resumen"]
    valores = {f[0].value: f[1].value for f in ws.iter_rows(min_row=4, max_row=6)}
    assert valores == {"Alumnos registrados": 4, "Terminaron todos los tests": 1, "En curso": 3}
    textos = [c.value for f in ws.iter_rows() for c in f if isinstance(c.value, str)]
    assert any("anterior" in t for t in textos)  # aviso de la versión vieja
    assert any(TEST1.titulo == t for t in textos) and any(TEST2.titulo == t for t in textos)


def test_hoja_alumnos_incluye_a_todos_con_su_estado(panel):
    ws = _libro(panel)["Alumnos"]
    cab = [c.value for c in ws[1]]
    assert cab[:8] == ["Matrícula", "Apellido paterno", "Apellido materno", "Nombre(s)", "Edad", "Tutor", "Registro", "Estado general"]
    filas = {f[0].value: f for f in ws.iter_rows(min_row=2)}
    assert set(filas) == {"100001", "100002", "100003", "100004"}
    beto = [c.value for c in filas["100002"]]
    assert "En curso" in beto and "10 de 25" in beto and "Terminado" in beto


def test_hoja_de_test_solo_trae_a_quien_tiene_resultado_con_porcentajes_numericos(panel):
    h1, h2 = _libro(panel)["Test 1"], _libro(panel)["Test 2"]
    assert {f[0].value for f in h1.iter_rows(min_row=2)} == {"100001", "100002"}
    assert {f[0].value for f in h2.iter_rows(min_row=2)} == {"100001"}
    primera = [c for c in h2[2]][7]
    assert isinstance(primera.value, (int, float)) and primera.number_format == "0.0%"


def test_hoja_respuestas_es_un_respaldo_completo():
    repo = poblar()
    repo.respuestas[("a1", TEST2.id)]["M01"] = None  # una "sin experiencia"
    panel = cargar_panel(repo, (TEST1, TEST2))
    ws = _libro(panel)["Respuestas"]
    filas = [tuple(c.value for c in f) for f in ws.iter_rows(min_row=2)]
    esperadas = sum(len(r) for t in panel.respuestas.values() for r in t.values())
    assert len(filas) == esperadas
    assert any(f[:3] == ("100001", TEST1.id, "BF01") for f in filas)
    assert ("100001", TEST2.id, "M01", "SE") in filas  # las "sin experiencia" se respaldan como SE


def test_el_concentrado_no_deja_pasar_formulas_ni_caracteres_ilegales():
    repo = poblar()
    repo.alumnos["100001"] = replace(repo.alumnos["100001"], nombres="=HYPERLINK(\"http://x\")", tutor="@cmd\x07x")
    p = cargar_panel(repo, (TEST1, TEST2))
    ws = openpyxl.load_workbook(io.BytesIO(excel_concentrado(p, FECHA)))["Alumnos"]
    celdas = [c.value for f in ws.iter_rows(min_row=2) for c in f if isinstance(c.value, str)]
    assert not any(v.startswith(("=", "@")) for v in celdas)
    assert any(v.startswith("'=") for v in celdas) and "'@cmdx" in celdas


def test_excel_concentrado_con_panel_vacio():
    vacio = PanelDatos((TEST1, TEST2), (), {TEST1.id: {}, TEST2.id: {}}, ())
    assert _libro(vacio).sheetnames[0] == "Resumen"


def test_zip_incluye_solo_a_quienes_terminaron_y_lista_los_omitidos(panel):
    datos, incluidos, omitidos = zip_pdfs(panel, FECHA)
    assert (incluidos, omitidos) == (1, 3)
    with zipfile.ZipFile(io.BytesIO(datos)) as z:
        assert z.testzip() is None
        nombres = z.namelist()
        assert sorted(nombres) == ["100001_Perez_Ana.pdf", "LEEME.txt"]
        assert z.read("100001_Perez_Ana.pdf").startswith(b"%PDF")
        leeme = z.read("LEEME.txt").decode()
        assert "Omitidos por no haber terminado todos los tests: 3" in leeme
        assert all(m in leeme for m in ("100002", "100003", "100004"))


def test_nombres_de_archivo_seguros():
    peligroso = Alumno("x", "../../etc/Ñandú", "L", "José/..\\Luis", "123456", 20, "T", "terminado")
    from modules.panel import FilaAlumno
    nombre = nombre_pdf(FilaAlumno(peligroso, (), (), ()))
    assert nombre == "123456_etc_Nandu_Jose_Luis.pdf" and "/" not in nombre and ".." not in nombre
    raro = replace(peligroso, apellido_paterno="李雷", nombres="Иван")
    assert nombre_pdf(FilaAlumno(raro, (), (), ())) == "123456_alumno_alumno.pdf"


def test_zip_sin_nadie_terminado_solo_trae_el_leeme():
    repo = poblar()
    repo.progreso.pop(("a1", TEST2.id))
    p = cargar_panel(repo, (TEST1, TEST2))
    datos, incluidos, omitidos = zip_pdfs(p, FECHA)
    assert incluidos == 0 and omitidos == 4
    with zipfile.ZipFile(io.BytesIO(datos)) as z:
        assert z.namelist() == ["LEEME.txt"]


def test_el_respaldo_incluye_respuestas_de_tests_y_preguntas_que_ya_no_estan_en_la_configuracion():
    repo = poblar()
    repo.respuestas[("a1", "test_retirado")] = {"R01": 3, "R02": None}              # test que ya no existe
    repo.respuestas[("a1", TEST1.id)]["BF_VIEJA"] = 4                              # pregunta que ya no existe
    panel = cargar_panel(repo, (TEST1, TEST2))
    ws = openpyxl.load_workbook(io.BytesIO(excel_concentrado(panel, FECHA)))["Respuestas"]
    filas = {tuple(c.value for c in f) for f in ws.iter_rows(min_row=2)}
    assert ("100001", "test_retirado", "R01", 3) in filas and ("100001", "test_retirado", "R02", "SE") in filas
    assert ("100001", TEST1.id, "BF_VIEJA", 4) in filas
    total = sum(len(r) for t in repo.obtener_respuestas_todas().values() for r in t.values())
    assert len(filas) == total  # nada se pierde


def test_panel_y_calificacion_ignoran_tests_retirados():
    repo = poblar()
    repo.respuestas[("a1", "test_retirado")] = {"R01": 3}
    panel = cargar_panel(repo, (TEST1, TEST2))
    assert panel.filas[0].completo and "test_retirado" in panel.respuestas
