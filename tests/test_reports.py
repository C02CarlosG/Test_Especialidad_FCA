import io
import json
import shutil
import subprocess
from datetime import datetime
from fractions import Fraction
from pathlib import Path

import openpyxl
import pytest

from fakes import alumno
from modules.config_loader import load_all_tests
from modules.reports import (FILAS_ENCABEZADO as OFF, excel_individual, grafica_radar, pct, pdf_individual,
                             vista_conjunta)
from modules.scoring import calificar

RAIZ = Path(__file__).resolve().parent.parent
TEST1, TEST2 = load_all_tests(RAIZ / "config" / "tests")
GOLDEN = json.loads((Path(__file__).parent / "golden" / "scoring_golden.json").read_text(encoding="utf-8"))
FECHA = datetime(2026, 10, 15)


def _resp(r):
    return {k: (None if v == "SE" else v) for k, v in r.items()}


@pytest.fixture(scope="module")
def resultados():
    return [calificar(TEST1, GOLDEN["test1"][40]["respuestas"]),
            calificar(TEST2, _resp(GOLDEN["test2"][20]["respuestas"]))]


def test_pct():
    assert pct(Fraction(3, 4)) == "75.0 %" and pct(None) == "Sin datos"


def test_vista_conjunta_comparte_especialidades_y_promedia(resultados):
    vc = vista_conjunta(resultados)
    assert {f.codigo for f in vc.filas} == {"MKT", "CH", "AF", "NI", "EMP"}
    for f in vc.filas:
        assert f.promedio == sum(f.porcentajes, Fraction(0)) / 2
    promedios = [f.promedio for f in vc.filas]
    assert promedios == sorted(promedios, reverse=True)
    assert vc.coinciden == (resultados[0].principal == resultados[1].principal)


def test_vista_conjunta_no_aplica_con_un_solo_test(resultados):
    assert vista_conjunta(resultados[:1]) is None
    assert grafica_radar(resultados[:1], ["x"]) is None


def test_vista_conjunta_promedio_requiere_dos_tests_con_dato():
    sin_dato = {k: ("SE" if k.startswith("M") else v) for k, v in GOLDEN["test2"][20]["respuestas"].items()}
    r = [calificar(TEST1, GOLDEN["test1"][0]["respuestas"]), calificar(TEST2, _resp(sin_dato))]
    mkt = next(f for f in vista_conjunta(r).filas if f.codigo == "MKT")
    assert mkt.porcentajes[1] is None and mkt.promedio is None
    assert vista_conjunta(r).filas[-1].codigo == "MKT"  # sin promedio va al final


def test_radar_es_un_png(resultados):
    png = grafica_radar(resultados, ["Test 1", "Test 2"])
    assert png.startswith(b"\x89PNG") and len(png) > 5000


def test_pdf_es_valido_y_contiene_los_datos(resultados):
    pdf = pdf_individual(alumno(), [TEST1, TEST2], resultados, FECHA)
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF")
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext no está instalado: solo se validó la estructura del PDF")
    texto = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pdf, capture_output=True, check=True).stdout.decode()
    for esperado in ("123456", "Ana Pérez López", "Test 1: The Big Five", "Vista conjunta", "15/10/2026",
                     "no son un diagnóstico", "Test 2: Administration Specialties"):
        assert esperado in texto
    for e in resultados[0].especialidades:
        assert e.nombre in texto


def test_pdf_con_un_test_sin_principal_no_falla():
    todo_se = {k: None for k in GOLDEN["test2"][0]["respuestas"]}
    r = calificar(TEST2, todo_se)
    assert r.principal is None
    pdf = pdf_individual(alumno(), [TEST2], [r], FECHA)
    assert pdf.startswith(b"%PDF")
    xlsx = excel_individual(alumno(), [TEST2], [r], FECHA)
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))["Test 1"]
    assert any(c.value == "Sin datos" for fila in ws.iter_rows() for c in fila)


def test_excel_hojas_y_valores(resultados):
    wb = openpyxl.load_workbook(io.BytesIO(excel_individual(alumno(), [TEST1, TEST2], resultados, FECHA)))
    assert wb.sheetnames == ["Resumen", "Test 1", "Test 2"]
    resumen = {fila[0].value: fila[1].value for fila in wb["Resumen"].iter_rows(min_row=OFF + 3, max_row=OFF + 7)}
    assert resumen["Matrícula"] == "123456" and resumen["Alumno"] == "Ana Pérez López"
    h1 = wb["Test 1"]
    filas = {f[0].value: f for f in h1.iter_rows(min_row=OFF + 4, max_row=OFF + 8)}
    r1 = resultados[0]
    for e in r1.especialidades:
        fila = filas[e.nombre]
        assert fila[1].value == e.obtenido and fila[2].value == e.maximo
        assert fila[3].value == pytest.approx(float(e.porcentaje)) and fila[3].number_format == "0.0%"
    assert wb["Test 2"].cell(OFF + 1, 1).value == TEST2.titulo


def test_excel_no_deja_pasar_formulas_del_alumno(resultados):
    from dataclasses import replace
    malo = replace(alumno(), nombres="=HYPERLINK(\"http://x\")", tutor="@cmd")
    wb = openpyxl.load_workbook(io.BytesIO(excel_individual(malo, [TEST1, TEST2], resultados, FECHA)))
    celdas = [c.value for fila in wb["Resumen"].iter_rows(min_row=OFF + 3, max_row=OFF + 7) for c in fila if isinstance(c.value, str)]
    assert not any(v.startswith(("=", "@")) for v in celdas)
    assert any(v.startswith("'=") for v in celdas) and any(v == "'@cmd" for v in celdas)


# ---- regresiones de la revisión independiente de la Fase 4 ----

@pytest.mark.parametrize("n,esperado", [(1, "6.3 %"), (5, "31.3 %"), (9, "56.3 %"), (13, "81.3 %")])
def test_pct_redondea_la_mitad_hacia_arriba_como_excel(n, esperado):
    # f"{x:.1f}" daría 6.2/31.2/56.2/81.2 (mitad a par) y el Excel mostraría .3
    assert pct(Fraction(n, 16)) == esperado


def test_pct_valores_basicos():
    assert [pct(Fraction(a, b)) for a, b in [(3, 4), (1, 3), (0, 1), (1, 1), (2, 3)]] == [
        "75.0 %", "33.3 %", "0.0 %", "100.0 %", "66.7 %"]


def test_pantalla_pdf_y_excel_muestran_el_mismo_porcentaje():
    todo = {k: 2 for k in GOLDEN["test2"][0]["respuestas"]}
    todo["M01"] = None  # una "SE": Mercadotecnia = 8/16 ... 9/16 con otro valor
    todo.update(M02=3, M04=2, M05=2, M09=2)
    r = calificar(TEST2, todo)
    mkt = r.especialidades[0]
    assert mkt.porcentaje == Fraction(9, 16)
    ws = openpyxl.load_workbook(io.BytesIO(excel_individual(alumno(), [TEST2], [r], FECHA)))["Test 1"]
    celda = next(f[3] for f in ws.iter_rows(min_row=OFF + 4, max_row=OFF + 8) if f[0].value == mkt.nombre)
    assert celda.value == 0.5625 and celda.number_format == "0.0%"  # Excel lo muestra 56.3 %
    assert pct(mkt.porcentaje) == "56.3 %"


def test_radar_no_dibuja_sin_datos_como_cero(resultados):
    from modules.reports import puntos_radar
    sin_mkt = {k: ("SE" if k.startswith("M") else v) for k, v in GOLDEN["test2"][20]["respuestas"].items()}
    r = [resultados[0], calificar(TEST2, _resp(sin_mkt))]
    vc = vista_conjunta(r)
    idx_mkt = next(k for k, f in enumerate(vc.filas) if f.codigo == "MKT")
    assert idx_mkt not in [k for k, _ in puntos_radar(vc, 1)]       # test 2: eje omitido
    assert idx_mkt in [k for k, _ in puntos_radar(vc, 0)]           # test 1 sí tiene dato
    assert all(0 < v <= 100 for _, v in puntos_radar(vc, 1))
    assert grafica_radar(r, ["T1", "T2"]).startswith(b"\x89PNG")


def test_radar_con_un_solo_eje_con_dato_no_falla(resultados):
    casi_todo_se = {k: (None if not k.startswith("M") else 3) for k in GOLDEN["test2"][0]["respuestas"]}
    r = [resultados[0], calificar(TEST2, casi_todo_se)]
    assert grafica_radar(r, ["T1", "T2"]).startswith(b"\x89PNG")


def test_el_radar_conserva_su_proporcion_en_el_pdf(resultados, tmp_path):
    from modules.reports import dimensiones_png
    png = grafica_radar(resultados, ["Test 1: The Big Five", "Test 2: Administration Specialties"])
    w, h = dimensiones_png(png)
    assert w > 300 and h > 300 and abs(w / h - 1.0) < 0.6
    if shutil.which("pdfimages"):
        (tmp_path / "r.pdf").write_bytes(pdf_individual(alumno(), [TEST1, TEST2], resultados, FECHA))
        salida = subprocess.run(["pdfimages", "-list", str(tmp_path / "r.pdf")], capture_output=True, text=True).stdout
        fila = [l for l in salida.splitlines() if " image " in l][0].split()
        x_ppi, y_ppi = float(fila[12]), float(fila[13])
        assert abs(x_ppi - y_ppi) / x_ppi < 0.02  # sin estirar


def test_el_pdf_imprime_nombres_fuera_de_windows_1252(resultados):
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext no está instalado")
    from dataclasses import replace
    raro = replace(alumno(), nombres="Łukasz Őrs ćšž", tutor="Dr. Ćirić")
    pdf = pdf_individual(raro, [TEST1, TEST2], resultados, FECHA)
    texto = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pdf, capture_output=True, check=True).stdout.decode()
    assert "Łukasz Őrs ćšž" in texto and "Dr. Ćirić" in texto and "■" not in texto


def test_excel_tolera_caracteres_de_control_aunque_la_validacion_falle(resultados):
    from dataclasses import replace
    malo = replace(alumno(), nombres="Ana\x01\x08Mar\x1bía", tutor="Mtro.\x07 Ruiz")
    wb = openpyxl.load_workbook(io.BytesIO(excel_individual(malo, [TEST1, TEST2], resultados, FECHA)))
    valores = [c.value for f in wb["Resumen"].iter_rows(min_row=OFF + 3, max_row=OFF + 7) for c in f]
    assert "AnaMaría Pérez López" in valores and "Mtro. Ruiz" in valores


def test_excel_individual_con_no_caracteres_sigue_siendo_valido(resultados):
    from dataclasses import replace
    malo = replace(alumno(), nombres="Ana\ufffe\uffff", tutor="Mtro.\ufffe Ruiz")
    xlsx = excel_individual(malo, [TEST1, TEST2], resultados, FECHA)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))  # antes: ParseError por XML mal formado
    valores = [c.value for f in wb["Resumen"].iter_rows(min_row=OFF + 3, max_row=OFF + 7) for c in f]
    assert "Ana Pérez López" in valores and "Mtro. Ruiz" in valores


def test_el_pdf_lleva_un_logo_a_cada_lado_del_titulo(resultados, tmp_path):
    if not shutil.which("pdfimages"):
        pytest.skip("pdfimages no está instalado")
    (tmp_path / "r.pdf").write_bytes(pdf_individual(alumno(), [TEST1, TEST2], resultados, FECHA))
    salida = subprocess.run(["pdfimages", "-list", "-f", "1", "-l", "1", str(tmp_path / "r.pdf")],
                            capture_output=True, text=True).stdout
    # cada logo es una imagen RGB más su máscara de transparencia (smask)
    filas = [l.split() for l in salida.splitlines()[2:] if l.strip()]  # columnas: página, n.º, tipo, ancho, alto...
    imagenes = [(f[3], f[4]) for f in filas if f[2] == "image"]
    assert imagenes == [("165", "240"), ("204", "240")]  # FCA a la izquierda, UABJO a la derecha
    assert sum(f[2] == "smask" for f in filas) == 2  # el fondo transparente se conserva


def test_el_excel_individual_lleva_los_logos_a_cada_lado_en_todas_las_hojas(resultados):
    wb = openpyxl.load_workbook(io.BytesIO(excel_individual(alumno(), [TEST1, TEST2], resultados, FECHA)))
    for ws in wb.worksheets:
        assert len(ws._images) == 2, ws.title
        izq, der = sorted(ws._images, key=lambda i: (i.anchor._from.col, i.anchor._from.colOff))
        assert izq.anchor._from.col == 0 and izq.anchor._from.row == 0           # FCA arriba a la izquierda
        assert der.anchor._from.col == (3 if ws.title == "Resumen" else 6)       # UABJO en la última columna usada
        px = lambda i: (round(i.anchor.ext.width / 9525), round(i.anchor.ext.height / 9525))  # noqa: E731
        assert px(izq) == (45, 66) and px(der)[1] == 66                          # misma altura, proporción conservada
        assert ws.cell(OFF + 1, 1).value  # el contenido bajó para no quedar debajo de los logos
        assert all(ws.cell(r, c).value is None for r in range(1, OFF + 1) for c in range(1, 8))  # banda libre
        assert ws.page_setup.orientation == "landscape" and ws.sheet_properties.pageSetUpPr.fitToPage


def test_el_excel_conserva_las_celdas_combinadas_y_alturas_tras_mover_el_contenido(resultados):
    ws = openpyxl.load_workbook(io.BytesIO(excel_individual(alumno(), [TEST1, TEST2], resultados, FECHA)))["Resumen"]
    unidas = [str(r) for r in ws.merged_cells.ranges]
    assert unidas == [f"A{9 + OFF}:D{9 + OFF}"]          # el aviso sigue combinado (A9:D9 antes del cambio)
    assert ws.row_dimensions[9 + OFF].height == 48 and ws.row_dimensions[9].height is None
    assert ws.cell(9 + OFF, 1).value.startswith("Estos resultados son una guía")
