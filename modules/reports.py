"""Reportes individuales del alumno (PDF y Excel) a partir de sus resultados calificados.

Todo se genera en memoria (bytes): no se escriben archivos ni se guardan en el servidor.
Los resultados son una guía de orientación, no un diagnóstico; los reportes lo dicen.

"Vista conjunta": las especialidades de la carrera que aparecen en todos los tests (mismo
código, p. ej. MKT) se comparan lado a lado y se promedian los porcentajes de los tests que
tienen dato. Es una lectura orientativa, no un puntaje oficial.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from fractions import Fraction
from pathlib import Path
from typing import Sequence
from xml.sax.saxutils import escape

from modules.acceso import Alumno
from modules.config_loader import TestConfig
from modules.scoring import ResultadoTest

LOGOS = (Path(__file__).resolve().parent.parent / "static" / "logo_fca.png",
         Path(__file__).resolve().parent.parent / "static" / "logo_uabjo.png")  # izquierda, derecha
AVISO = ("Estos resultados son una guía de orientación para elegir tu especialidad; no son un "
         "diagnóstico psicológico ni una decisión obligatoria. Coméntalos con tu tutor u orientador.")
COLORES = ("#0F1E4A", "#E5B300", "#17683F", "#B23A48")  # marino, oro, verde, granate: un color por test (hasta 4)


@dataclass(frozen=True)
class FilaConjunta:
    codigo: str
    nombre: str
    porcentajes: tuple[Fraction | None, ...]  # uno por test, en el orden de los tests
    promedio: Fraction | None  # None si menos de 2 tests tienen dato


@dataclass(frozen=True)
class VistaConjunta:
    filas: tuple[FilaConjunta, ...]  # de mayor a menor promedio; sin promedio al final
    coinciden: bool  # todos los tests señalan la misma especialidad principal


def _seguro(texto: str) -> str:
    """Texto del alumno listo para una celda de Excel: sin caracteres que openpyxl rechaza y sin
    que Excel lo interprete como fórmula (=, +, -, @)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from modules.security import limpiar_caracteres
    texto = ILLEGAL_CHARACTERS_RE.sub("", limpiar_caracteres(texto))
    return "'" + texto if texto.startswith(("=", "+", "-", "@", "\t", "\r")) else texto


def pct(valor: Fraction | None) -> str:
    """Porcentaje con un decimal, redondeando .05 hacia arriba (igual que el formato 0.0 % de Excel;
    f"{x:.1f}" redondea a par: 56.25 daría 56.2 en pantalla y 56.3 en el Excel)."""
    if valor is None:
        return "Sin datos"
    decimas = int(valor * 1000 + Fraction(1, 2))  # valor >= 0: int() equivale a floor
    return f"{decimas // 10}.{decimas % 10} %"


def vista_conjunta(resultados: Sequence[ResultadoTest]) -> VistaConjunta | None:
    """None si hay un solo test o ninguna especialidad común a todos."""
    if len(resultados) < 2:
        return None
    por_test = [{a.codigo: a for a in r.afinidades} for r in resultados]
    comunes = [a.codigo for a in resultados[0].afinidades if all(a.codigo in d for d in por_test)]
    if not comunes:
        return None
    filas = []
    for cod in comunes:
        porcentajes = tuple(d[cod].porcentaje for d in por_test)
        con_dato = [p for p in porcentajes if p is not None]
        promedio = sum(con_dato, Fraction(0)) / len(con_dato) if len(con_dato) >= 2 else None
        filas.append(FilaConjunta(cod, por_test[0][cod].nombre, porcentajes, promedio))
    filas.sort(key=lambda f: (f.promedio is not None, f.promedio or Fraction(0)), reverse=True)
    principales = {r.principal for r in resultados}
    return VistaConjunta(tuple(filas), len(principales) == 1 and None not in principales)


def _nombre_de(r: ResultadoTest, codigo: str | None) -> str:
    if codigo is None:
        return "Sin datos suficientes"
    return next(a.nombre for a in r.afinidades if a.codigo == codigo)


def _descripcion(test: TestConfig, codigo: str | None) -> str:
    """Texto de la especialidad: afinidad derivada (test 1) o descripción de la especialidad."""
    if codigo is None:
        return ""
    for d in test.afinidades:
        if d.codigo == codigo:
            return d.descripcion
    esp = next((e for e in test.especialidades if e.codigo == codigo), None)
    return " ".join(x for x in (esp.descripcion, esp.perfil_profesional) if x) if esp else ""


def _notas(r: ResultadoTest) -> list[str]:
    notas = []
    if r.empate_tecnico:
        notas.append("Empate técnico: tus dos primeras opciones están a 5 puntos porcentuales o menos; "
                     "considera ambas.")
    for e in r.especialidades:
        if e.sin_datos:
            notas.append(f"{e.nombre}: contestaste «sin experiencia» en todo, no hay dato para compararla.")
        elif e.requiere_exploracion:
            notas.append(f"{e.nombre}: contestaste «sin experiencia» en {e.sin_experiencia} reactivos; "
                         "conviene explorar esta área antes de descartarla.")
    if r.principal is None:
        notas.append("No hubo datos suficientes para sugerir una primera y segunda opción en este test.")
    return notas


# --------------------------------------------------------------------------- gráfica

def puntos_radar(vc: VistaConjunta, i: int) -> list[tuple[int, float]]:
    """(índice del eje, porcentaje) del test `i`, solo de las especialidades con dato.
    "Sin datos" no se dibuja como 0 %: se omite el eje."""
    return [(k, float(f.porcentajes[i]) * 100) for k, f in enumerate(vc.filas) if f.porcentajes[i] is not None]


def dimensiones_png(png: bytes) -> tuple[int, int]:
    """(ancho, alto) en píxeles leídos de la cabecera IHDR del PNG."""
    import struct
    return struct.unpack(">II", png[16:24])


def grafica_radar(resultados: Sequence[ResultadoTest], titulos: Sequence[str]) -> bytes | None:
    """Radar con el porcentaje de cada test por especialidad común. None si no aplica."""
    vc = vista_conjunta(resultados)
    if vc is None or len(vc.filas) < 3:
        return None
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from math import pi

    etiquetas = [f.nombre for f in vc.filas]
    n = len(etiquetas)
    angulos = [i / n * 2 * pi for i in range(n)] + [0]
    fig, ax = plt.subplots(figsize=(6.2, 5.2), subplot_kw={"projection": "polar"})
    for i, titulo in enumerate(titulos):
        puntos = puntos_radar(vc, i)
        color = COLORES[i % len(COLORES)]
        if len(puntos) < 2:  # con un solo eje no hay figura que dibujar
            ax.plot([], [], color=color, linewidth=2, label=f"{titulo} (sin datos suficientes)")
            continue
        ang = [angulos[k] for k, _ in puntos] + [angulos[puntos[0][0]]]
        val = [v for _, v in puntos] + [puntos[0][1]]
        ax.plot(ang, val, color=color, linewidth=2, label=titulo)
        ax.fill(ang, val, color=color, alpha=0.12)
    ax.set_xticks(angulos[:-1])
    ax.set_xticklabels([e.replace(" ", "\n", 1) if len(e) > 14 else e for e in etiquetas], fontsize=8)
    ax.set_ylim(0, 100)
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(["20 %", "40 %", "60 %", "80 %", "100 %"], fontsize=7, color="#666")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), fontsize=8, ncol=2, frameon=False)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF

def _registrar_fuentes() -> tuple[str, str]:
    """(normal, negrita). DejaVu Sans viene con matplotlib y cubre acentos y alfabetos latino,
    cirílico y griego; Helvetica (la de reportlab) imprime '■' fuera de Windows-1252."""
    try:
        import matplotlib
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        carpeta = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(carpeta / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(carpeta / "DejaVuSans-Bold.ttf")))
        return "DejaVuSans", "DejaVuSans-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


FUENTES_PDF = Path(__file__).resolve().parent.parent / "static" / "fonts" / "pdf"


def _fuentes_marca() -> tuple[str, str, str] | None:
    """(texto, texto en negrita, títulos) con la tipografía de la app: Plus Jakarta Sans y Fraunces.
    None si faltan los archivos: el PDF usa entonces DejaVu Sans."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        archivos = {"Jakarta": "plus-jakarta-sans-latin-400-normal.ttf", "Jakarta-Bold": "plus-jakarta-sans-latin-700-normal.ttf",
                    "Fraunces": "fraunces-latin-600-normal.ttf"}
        for nombre, archivo in archivos.items():
            if nombre not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(nombre, str(FUENTES_PDF / archivo)))
        return "Jakarta", "Jakarta-Bold", "Fraunces"
    except Exception:
        return None


def _con_respaldo(texto: str, fuente: str, respaldo: str) -> str:
    """Escapa `texto` para un Paragraph y pone en la fuente de respaldo los caracteres que `fuente`
    no tiene (alfabetos no latinos, 'ł', 'ő'...), para que nunca salga '■'."""
    from reportlab.pdfbase import pdfmetrics
    cubiertos = pdfmetrics.getFont(fuente).face.charToGlyph
    salida, fuera = [], False
    for ch in texto:
        sin = ord(ch) not in cubiertos and not ch.isspace()
        if sin != fuera:
            salida.append(f'<font name="{respaldo}">' if sin else "</font>")
            fuera = sin
        salida.append(escape(ch))
    if fuera:
        salida.append("</font>")
    return "".join(salida)


def pdf_individual(alumno: Alumno, tests: Sequence[TestConfig], resultados: Sequence[ResultadoTest],
                   fecha: datetime) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    azul = colors.HexColor(COLORES[0])
    base = getSampleStyleSheet()
    respaldo, respaldo_neg = _registrar_fuentes()  # DejaVu: cubre cualquier alfabeto
    marca = _fuentes_marca()
    normal, negrita, titulos = marca if marca else (respaldo, respaldo_neg, respaldo_neg)
    tinta = colors.HexColor("#1E2A4A")
    h1 = ParagraphStyle("h1", parent=base["Title"], fontName=titulos, fontSize=20, leading=24, textColor=azul, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName=titulos, fontSize=14, textColor=azul, spaceBefore=14,
                        spaceAfter=5)
    h3 = ParagraphStyle("h3", parent=base["Heading3"], fontName=negrita, fontSize=8, leading=10,
                        textColor=colors.HexColor("#5A6788"))
    cuerpo = ParagraphStyle("c", parent=base["BodyText"], fontName=normal, fontSize=9.5, leading=13.5, textColor=tinta)
    chico = ParagraphStyle("s", parent=cuerpo, fontSize=8.5, textColor=colors.HexColor("#5A6788"))
    celda = ParagraphStyle("t", parent=cuerpo, fontSize=8.3, leading=11)
    celda_cab = ParagraphStyle("tc", parent=celda, fontName=negrita, textColor=colors.white)

    def _txt(texto, estilo):
        return _con_respaldo(str(texto), estilo.fontName, respaldo)

    def p(texto, estilo=cuerpo):
        return Paragraph(_txt(texto, estilo), estilo)

    def tabla(filas, anchos, cabecera=True):
        datos = [[Paragraph(_txt(c, celda_cab if cabecera and i == 0 else celda), celda_cab if cabecera and i == 0 else celda)
                  for c in fila] for i, fila in enumerate(filas)]
        t = Table(datos, colWidths=anchos, repeatRows=1 if cabecera else 0)
        estilo = [("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6DDEB")),
                  ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                  ("TOPPADDING", (0, 0), (-1, -1), 4.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5)]
        if cabecera:
            estilo += [("BACKGROUND", (0, 0), (-1, 0), azul),
                       ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FC")])]
        else:
            estilo += [("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F4F7FC"))]
        t.setStyle(TableStyle(estilo))
        return t

    ancho = letter[0] - 4 * cm

    def _encabezado_pdf(p, h1, h3, ancho):
        """Logo de la FCA a la izquierda, título al centro y logo de la UABJO a la derecha."""
        alto = 2.3 * cm

        def logo(ruta):
            if not ruta.is_file():
                return ""
            from PIL import Image as PILImage
            w, h = PILImage.open(ruta).size
            return Image(str(ruta), width=alto * w / h, height=alto, mask="auto")

        centro = [Paragraph("UNIVERSIDAD AUTÓNOMA BENITO JUÁREZ DE OAXACA", ParagraphStyle("k", parent=h3, alignment=1, fontSize=6.6, spaceAfter=3)),
                  Paragraph("Evaluación Diagnóstica FCA", ParagraphStyle("h1c", parent=h1, alignment=1)),
                  Paragraph("Reporte de resultados", ParagraphStyle("h3c", parent=h3, alignment=1, fontSize=8.5))]
        t = Table([[logo(LOGOS[0]), centro, logo(LOGOS[1])]], colWidths=[2.9 * cm, ancho - 5.8 * cm, 2.9 * cm])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "LEFT"),
                               ("ALIGN", (2, 0), (2, 0), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.HexColor("#E5B300")),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        return t
    historia = [_encabezado_pdf(p, h1, h3, ancho),
                tabla([["Alumno", alumno.nombre_completo], ["Matrícula", alumno.matricula],
                       ["Edad", str(alumno.edad)], ["Tutor / orientador", alumno.tutor],
                       ["Fecha", fecha.strftime("%d/%m/%Y")]], [4 * cm, ancho - 4 * cm], cabecera=False),
                Spacer(1, 6), p(AVISO, chico)]

    historia.append(p("Resumen", h2))
    filas = [["Test", "1ª opción", "2ª opción"]]
    for t, r in zip(tests, resultados):
        filas.append([t.titulo, _nombre_de(r, r.principal), _nombre_de(r, r.secundaria)])
    historia.append(tabla(filas, [ancho * 0.4, ancho * 0.3, ancho * 0.3]))

    vc = vista_conjunta(resultados)
    if vc is not None:
        historia.append(p("Vista conjunta de los tests", h2))
        if vc.coinciden:
            historia.append(p("Todos los tests apuntan a la misma primera opción: "
                              f"{_nombre_de(resultados[0], resultados[0].principal)}."))
        else:
            historia.append(p("Los tests no apuntan a la misma primera opción; compara las columnas y "
                              "coméntalo con tu tutor."))
        cab = ["Especialidad"] + [t.titulo for t in tests] + ["Promedio"]
        filas = [cab] + [[f.nombre, *[pct(x) for x in f.porcentajes], pct(f.promedio)] for f in vc.filas]
        resto = ancho - ancho * 0.28
        historia.append(tabla(filas, [ancho * 0.28] + [resto * 0.72 / len(tests)] * len(tests) + [resto * 0.28]))
        historia.append(p("El promedio solo considera los tests con dato y es orientativo.", chico))
        png = grafica_radar(resultados, [t.titulo for t in tests])
        if png:
            w_px, h_px = dimensiones_png(png)
            historia += [Spacer(1, 6), Image(io.BytesIO(png), width=12.5 * cm, height=12.5 * cm * h_px / w_px)]

    for t, r in zip(tests, resultados):
        historia.append(p(t.titulo, h2))
        filas = [["Área", "Puntaje", "Porcentaje", "Nivel"]]
        for e in r.especialidades:
            filas.append([e.nombre, "—" if e.sin_datos else f"{e.obtenido} de {e.maximo}", pct(e.porcentaje),
                          e.nivel or "—"])
        historia.append(tabla(filas, [ancho * 0.4, ancho * 0.2, ancho * 0.18, ancho * 0.22]))
        if t.afinidades:
            historia += [Spacer(1, 6), p("Afinidad con las especialidades de la carrera", cuerpo)]
            filas = [["Especialidad", "Porcentaje", "Nivel"]]
            for a in r.afinidades:
                filas.append([a.nombre, pct(a.porcentaje), a.nivel or "—"])
            historia.append(tabla(filas, [ancho * 0.45, ancho * 0.25, ancho * 0.3]))
        for etiqueta, cod in (("1ª opción", r.principal), ("2ª opción", r.secundaria)):
            if cod is None:
                continue
            historia += [Spacer(1, 4), p(f"{etiqueta}: {_nombre_de(r, cod)}", cuerpo)]
            if desc := _descripcion(t, cod):
                historia.append(p(desc, chico))
        for nota in _notas(r):
            historia.append(p("• " + nota, chico))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=f"Resultados {alumno.matricula}",
                      author="Evaluación Diagnóstica FCA").build(historia)
    return buf.getvalue()


# --------------------------------------------------------------------------- Excel

FILAS_ENCABEZADO = 5  # filas que quedan arriba, libres, para los logos
_ALTO_LOGO_PX = 66


def con_logos(ws, col_derecha: int) -> None:
    """Reserva FILAS_ENCABEZADO filas arriba (mueve el contenido hacia abajo) y pone el logo de la FCA a
    la izquierda y el de la UABJO alineado al borde derecho de la columna `col_derecha` (1 = A).
    Llamar al final, con los anchos de columna ya definidos."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Border, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    unidos = [r.coord for r in ws.merged_cells.ranges]
    alturas = {i: d.height for i, d in ws.row_dimensions.items() if d.height}
    for coord in unidos:
        ws.unmerge_cells(coord)
    if ws.max_row > 1 or ws.cell(1, 1).value is not None:
        ws.move_range(ws.dimensions, rows=FILAS_ENCABEZADO)
    for coord in unidos:
        from openpyxl.worksheet.cell_range import CellRange
        r = CellRange(coord)
        r.shift(row_shift=FILAS_ENCABEZADO)
        ws.merge_cells(r.coord)
    for i in alturas:
        ws.row_dimensions[i].height = None
    for i, h in alturas.items():
        ws.row_dimensions[i + FILAS_ENCABEZADO].height = h
    for i in range(1, FILAS_ENCABEZADO):
        ws.row_dimensions[i].height = 14  # 4 filas x 14 pt: caben los logos
    ws.row_dimensions[FILAS_ENCABEZADO].height = 8
    linea = Border(bottom=Side(style="medium", color="E5B300"))
    for c in range(1, col_derecha + 1):
        ws.cell(FILAS_ENCABEZADO - 1, c).border = linea

    def poner(ruta: Path, col: int, derecha: bool) -> None:
        if not ruta.is_file():
            return
        img = XLImage(str(ruta))
        img.width, img.height = round(img.width * _ALTO_LOGO_PX / img.height), _ALTO_LOGO_PX
        ancho_col = int((ws.column_dimensions[get_column_letter(col)].width or 13) * 7 + 5)
        desplazo = max(0, ancho_col - img.width - 4) if derecha else 4
        ancla = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(desplazo), row=0, rowOff=pixels_to_EMU(4))
        img.anchor = OneCellAnchor(_from=ancla, ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)))
        ws.add_image(img)

    poner(LOGOS[0], 1, derecha=False)
    poner(LOGOS[1], col_derecha, derecha=True)
    # Al imprimir o exportar a PDF, que quepa a lo ancho de una página (si no, el logo derecho cae en otra).
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_setup.orientation = "landscape"

def excel_individual(alumno: Alumno, tests: Sequence[TestConfig], resultados: Sequence[ResultadoTest],
                     fecha: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    negrita, cab = Font(bold=True), PatternFill("solid", fgColor="E8EDF8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    def fila_cab(hoja, fila, valores):
        for c, v in enumerate(valores, start=1):
            celda = hoja.cell(fila, c, v)
            celda.font, celda.fill = negrita, cab

    def porcentaje(hoja, fila, col, valor):
        c = hoja.cell(fila, col, None if valor is None else float(valor))
        c.number_format = "0.0%"
        if valor is None:
            c.value = "Sin datos"
        return c

    ws["A1"] = "Evaluación Diagnóstica FCA - Reporte de resultados"
    ws["A1"].font = Font(bold=True, size=14)
    datos = [("Alumno", _seguro(alumno.nombre_completo)), ("Matrícula", alumno.matricula),
             ("Edad", alumno.edad), ("Tutor / orientador", _seguro(alumno.tutor)),
             ("Fecha", fecha.strftime("%d/%m/%Y"))]
    for i, (k, v) in enumerate(datos, start=3):
        ws.cell(i, 1, k).font = negrita
        ws.cell(i, 2, v)
    ws.cell(9, 1, AVISO).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A9:D9")
    ws.row_dimensions[9].height = 48

    fila = 11
    fila_cab(ws, fila, ["Test", "1ª opción", "2ª opción", "Empate técnico"])
    for t, r in zip(tests, resultados):
        fila += 1
        for c, v in enumerate([t.titulo, _nombre_de(r, r.principal), _nombre_de(r, r.secundaria),
                               "Sí" if r.empate_tecnico else "No"], start=1):
            ws.cell(fila, c, v)

    vc = vista_conjunta(resultados)
    if vc is not None:
        fila += 2
        ws.cell(fila, 1, "Vista conjunta (el promedio es orientativo)").font = negrita
        fila += 1
        fila_cab(ws, fila, ["Especialidad", *[t.titulo for t in tests], "Promedio"])
        for f in vc.filas:
            fila += 1
            ws.cell(fila, 1, f.nombre)
            for i, x in enumerate(f.porcentajes, start=2):
                porcentaje(ws, fila, i, x)
            porcentaje(ws, fila, 2 + len(tests), f.promedio)

    for n, (t, r) in enumerate(zip(tests, resultados), start=1):
        h = wb.create_sheet(f"Test {n}")
        h["A1"] = t.titulo
        h["A1"].font = Font(bold=True, size=13)
        fila_cab(h, 3, ["Área", "Puntaje obtenido", "Puntaje máximo", "Porcentaje", "Nivel", "Respuestas «sin experiencia»", "Requiere explorar"])
        fila = 3
        for e in r.especialidades:
            fila += 1
            h.cell(fila, 1, e.nombre)
            h.cell(fila, 2, None if e.sin_datos else e.obtenido)
            h.cell(fila, 3, None if e.sin_datos else e.maximo)
            porcentaje(h, fila, 4, e.porcentaje)
            h.cell(fila, 5, e.nivel or "—")
            h.cell(fila, 6, e.sin_experiencia)
            h.cell(fila, 7, "Sí" if e.requiere_exploracion else "No")
        if t.afinidades:
            fila += 2
            h.cell(fila, 1, "Afinidad con las especialidades de la carrera").font = negrita
            fila += 1
            fila_cab(h, fila, ["Especialidad", "Puntaje obtenido", "Puntaje máximo", "Porcentaje", "Nivel"])
            for a in r.afinidades:
                fila += 1
                h.cell(fila, 1, a.nombre)
                h.cell(fila, 2, None if a.sin_datos else a.obtenido)
                h.cell(fila, 3, None if a.sin_datos else a.maximo)
                porcentaje(h, fila, 4, a.porcentaje)
                h.cell(fila, 5, a.nivel or "—")
        fila += 2
        for etiqueta, cod in (("1ª opción", r.principal), ("2ª opción", r.secundaria)):
            h.cell(fila, 1, etiqueta).font = negrita
            h.cell(fila, 2, _nombre_de(r, cod))
            fila += 1
        for nota in _notas(r):
            fila += 1
            h.cell(fila, 1, nota)

    for hoja in wb.worksheets:
        for col in range(1, 8):
            hoja.column_dimensions[get_column_letter(col)].width = 34 if col == 1 else 26
    for hoja in wb.worksheets:
        con_logos(hoja, 4 if hoja.title == "Resumen" else 7)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
